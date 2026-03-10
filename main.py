import subprocess
import time
import threading
import signal
import shutil
from pathlib import Path
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from queue import Queue

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ==================== CONFIG ====================
BASE_DIR = Path(__file__).resolve().parent
DEVICES_FILE = BASE_DIR / "devices.txt"
EXCEL_FILE = BASE_DIR / "battery_log.xlsx"
CSV_BACKUP_FILE = BASE_DIR / "battery_log_backup.csv"

POLL_INTERVAL_SEC = 5
ADB_TIMEOUT_SEC = 60
MAX_WORKERS = 12

ADB_PATH = BASE_DIR / "platform-tools" / "adb.exe"

# ==================== GLOBAL STATE ====================
TEST_START_TIME = None
CAPACITY_PER_DEVICE = {}
_excel_write_lock = threading.Lock()
_write_queue = Queue()
_writer_thread = None
_shutdown_flag = False


def ensure_adb_server_running() -> None:
    """Перезапускает ADB сервер для чистого состояния"""
    print("🔄 Перезапуск ADB сервера...")
    subprocess.run(
        [str(ADB_PATH), "kill-server"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL
    )
    time.sleep(5)
    subprocess.run(
        [str(ADB_PATH), "start-server"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL
    )
    time.sleep(5)
    print("✅ ADB сервер запущен")


def connect_to_all_devices(devices: list[str]) -> None:
    """Подключается ко всем устройствам по списку"""
    print("\n--- 🔗 Подключение к устройствам ---")
    connected = 0
    for device in devices:
        print(f"  Connecting to {device}...")
        result = subprocess.run(
            [str(ADB_PATH), "connect", device],
            capture_output=True,
            text=True,
            timeout=30
        )
        if result.returncode == 0 and "connected" in result.stdout.lower():
            print(f"  ✓ Connected to {device}")
            connected += 1
        else:
            err = result.stderr.strip() or result.stdout.strip()
            print(f"  ✗ Failed to connect to {device}: {err}")
    print(f"--- Подключено: {connected}/{len(devices)} ---\n")


def load_devices() -> list[str]:
    """Загружает список устройств из файла"""
    if not DEVICES_FILE.exists():
        raise FileNotFoundError(f"Файл {DEVICES_FILE} не найден")

    devices = []
    with DEVICES_FILE.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            devices.append(line)
    return devices


def is_device_connected(device: str) -> bool:
    """Проверяет, доступно ли устройство через ADB"""
    try:
        result = subprocess.run(
            [str(ADB_PATH), "-s", device, "shell", "echo", "ok"],
            capture_output=True,
            text=True,
            timeout=ADB_TIMEOUT_SEC
        )
        if result.returncode != 0:
            return False
        return result.stdout.strip() == "ok"
    except Exception:
        return False


def adb_read(device: str, path: str) -> int | None:
    """Читает числовое значение из системного файла на устройстве"""
    try:
        result = subprocess.run(
            [str(ADB_PATH), "-s", device, "shell", "cat", path],
            capture_output=True,
            text=True,
            timeout=ADB_TIMEOUT_SEC
        )
        if result.returncode != 0:
            return None
        value_str = result.stdout.strip()
        if not value_str:
            return None
        return int(value_str)
    except (subprocess.TimeoutExpired, ValueError, Exception):
        return None


def poll_device(device: str) -> dict:
    """Опрашивает устройство и возвращает словарь с метриками"""
    if not is_device_connected(device):
        return {
            "device": device,
            "voltage_mv": None,
            "current_ma": None,
            "cpu_temp1": None,
            "battery_temp1": None,
            "battery_temp2": None,
            "status": "OFFLINE"
        }

    voltage_raw = adb_read(
        device,
        "/sys/class/power_supply/battery/voltage_now"
    )
    current_raw = adb_read(
        device,
        "/sys/class/power_supply/battery/current_now"
    )
    cpu_temp1 = adb_read(
        device,
        "/sys/class/thermal/thermal_zone5/temp"
    )
    battery_temp1 = adb_read(
        device,
        "/sys/class/thermal/thermal_zone27/temp"
    )
    battery_temp2 = adb_read(
        device,
        "/sys/class/thermal/thermal_zone28/temp"
    )

    if voltage_raw is None or current_raw is None:
        return {
            "device": device,
            "voltage_mv": None,
            "current_ma": None,
            "cpu_temp1": None,
            "battery_temp1": None,
            "battery_temp2": None,
            "status": "NO_DATA"
        }

    return {
        "device": device,
        "voltage_mv": voltage_raw // 1000,
        "current_ma": current_raw // 1000,
        "cpu_temp1": cpu_temp1 // 1000 if cpu_temp1 else None,
        "battery_temp1": battery_temp1 // 1000 if battery_temp1 else None,
        "battery_temp2": battery_temp2 // 1000 if battery_temp2 else None,
        "status": "OK"
    }


def _backup_to_csv(device: str, row_data: list) -> None:
    """Экстренное сохранение строки в CSV при сбое Excel"""
    try:
        with open(CSV_BACKUP_FILE, "a", encoding="utf-8-sig") as f:
            f.write(",".join(str(v) for v in row_data) + "\n")
    except Exception as e:
        print(f"[CRITICAL] CSV backup failed for {device}: {e}")


def _write_row_to_excel(
        device: str,
        row_data: list,
        timestamp: datetime
) -> bool:
    """Внутренняя функция записи одной строки в Excel (вызывается внутри lock)"""
    global CAPACITY_PER_DEVICE

    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        sheet_name = device.replace(":", "_")

        if sheet_name not in wb.sheetnames:
            print(f"[WARN] Sheet {sheet_name} not found, skipping write")
            return False

        ws = wb[sheet_name]

        # Добавляем Test Duration
        test_duration = int((timestamp - TEST_START_TIME).total_seconds())
        row_data.append(test_duration)

        # Рассчитываем Capacity (mAh)
        current_ma = row_data[3]
        if current_ma is not None:
            delta_capacity = abs(current_ma) * POLL_INTERVAL_SEC / 3600
            CAPACITY_PER_DEVICE[device] = (
                    CAPACITY_PER_DEVICE.get(device, 0.0) + delta_capacity
            )
        row_data.append(round(CAPACITY_PER_DEVICE.get(device, 0.0), 4))

        # Добавляем строку
        ws.append(row_data)

        # Автоширина колонок (только первые 100 строк)
        if ws.max_row <= 101:
            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = min(
                    max_length + 2,
                    50
                )

        # Цветовая индикация статуса
        status = row_data[7]
        fill_color = None
        if status == "OK":
            fill_color = PatternFill(
                start_color="D5F5E3",
                end_color="D5F5E3",
                fill_type="solid"
            )
        elif status == "OFFLINE":
            fill_color = PatternFill(
                start_color="FADBD8",
                end_color="FADBD8",
                fill_type="solid"
            )
        elif status == "NO_DATA":
            fill_color = PatternFill(
                start_color="FEF9E7",
                end_color="FEF9E7",
                fill_type="solid"
            )

        if fill_color:
            for col_num in range(1, len(row_data) + 1):
                cell = ws.cell(row=ws.max_row, column=col_num)
                cell.fill = fill_color

        wb.save(EXCEL_FILE)
        return True

    except Exception as e:
        print(f"[ERROR] Excel write failed for {device}: {e}")
        return False


def _excel_writer_worker() -> None:
    """Фоновый поток для безопасной записи в Excel"""
    global _shutdown_flag

    print("📝 Writer thread started")

    while not _shutdown_flag or not _write_queue.empty():
        try:
            device, row_data, timestamp = _write_queue.get(timeout=1.0)

            with _excel_write_lock:
                success = _write_row_to_excel(device, row_data, timestamp)
                if not success:
                    _backup_to_csv(device, row_data)

            _write_queue.task_done()

        except Exception as e:
            print(f"[ERROR] Writer thread: {e}")
            time.sleep(0.5)

    print("📝 Writer thread stopped")


def append_to_excel_async(device: str, row_data: list) -> None:
    """Публичный API: добавляет данные в очередь записи (не блокирует)"""
    _write_queue.put((device, row_data.copy(), datetime.now()))


def init_excel(devices: list[str]) -> None:
    """Инициализирует Excel файл с листами для каждого устройства"""
    global TEST_START_TIME, CAPACITY_PER_DEVICE

    TEST_START_TIME = datetime.now()

    for device in devices:
        CAPACITY_PER_DEVICE[device] = 0.0

    # Проверяем шаблон
    template = BASE_DIR / "battery_log_template.xlsx"
    if template.exists():
        shutil.copy(template, EXCEL_FILE)
        print(f"📋 Использован шаблон: {template.name}")
        return

    # Создаём новый файл
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    headers = [
        "Timestamp",
        "Device",
        "Voltage (mV)",
        "Current (mA)",
        "CPU temp(cpu_big1)",
        "Battery Temp #1",
        "Battery Temp #2",
        "Status",
        "Test Duration (sec)",
        "Capacity (mAh)"
    ]

    for device in devices:
        safe_name = device.replace(":", "_").replace("/", "_").replace("\\", "_")
        safe_name = safe_name[:31]

        ws = wb.create_sheet(title=safe_name)

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="4472C4",
                end_color="4472C4",
                fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(col_num)].width = (
                    len(header) + 2
            )

    wb.save(EXCEL_FILE)
    print(f"✅ Excel файл создан: {EXCEL_FILE.name}")


def _graceful_shutdown(signum=None, frame=None) -> None:
    """Обработчик корректного завершения"""
    global _shutdown_flag

    print("\n⏹ Получен сигнал остановки...")
    _shutdown_flag = True

    print("⏳ Ожидание завершения записи данных...")
    _write_queue.join()

    if _writer_thread and _writer_thread.is_alive():
        _writer_thread.join(timeout=5.0)

    print("✅ Все данные сохранены. Можно закрывать программу.")


def main() -> None:
    global MAX_WORKERS, _writer_thread

    print(f"🔋 Battery Monitor v1.2")
    print(f"📁 Working directory: {BASE_DIR}")
    print(f"🔧 ADB path: {ADB_PATH} (exists: {ADB_PATH.exists()})\n")

    if not ADB_PATH.exists():
        print(f"❌ ADB не найден: {ADB_PATH}")
        print(
            "💡 Скачайте platform-tools: "
            "https://developer.android.com/tools/releases/platform-tools"
        )
        return

    ensure_adb_server_running()

    try:
        devices = load_devices()
    except FileNotFoundError as e:
        print(f"❌ {e}")
        return

    if not devices:
        print("❌ Список устройств пуст")
        return

    MAX_WORKERS = min(12, len(devices))
    print(f"🔧 Workers: {MAX_WORKERS} (устройств: {len(devices)})")

    connect_to_all_devices(devices)

    init_excel(devices)

    _writer_thread = threading.Thread(
        target=_excel_writer_worker,
        daemon=True
    )
    _writer_thread.start()

    signal.signal(signal.SIGINT, _graceful_shutdown)
    signal.signal(signal.SIGTERM, _graceful_shutdown)

    print(f"\n🚀 Старт мониторинга ({len(devices)} устройств, интервал {POLL_INTERVAL_SEC}с)")
    print("💡 Нажмите Ctrl+C для корректной остановки\n")

    iteration = 0
    start_time = datetime.now()

    try:
        while not _shutdown_flag:
            iteration += 1
            timestamp = datetime.now().isoformat(timespec="seconds")

            with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
                futures = {
                    executor.submit(poll_device, d): d
                    for d in devices
                }

                for future in as_completed(futures):
                    device = futures[future]
                    try:
                        result = future.result(timeout=ADB_TIMEOUT_SEC)

                        row = [
                            timestamp,
                            result["device"],
                            result["voltage_mv"],
                            result["current_ma"],
                            result["cpu_temp1"],
                            result["battery_temp1"],
                            result["battery_temp2"],
                            result["status"],
                        ]

                        print(
                            f"[{result['device']}] "
                            f"{result['status']:8s} "
                            f"V={result['voltage_mv']:5d}mV "
                            f"I={result['current_ma']:6d}mA "
                            f"CPU={result['cpu_temp1']:3d}°C "
                            f"BAT1={result['battery_temp1']:3d}°C "
                            f"BAT2={result['battery_temp2']:3d}°C"
                        )

                        append_to_excel_async(result["device"], row)

                    except Exception as e:
                        print(f"[ERROR] Polling {device}: {e}")

            if iteration % 12 == 0:
                elapsed = (datetime.now() - start_time).total_seconds() / 60
                print(f"⏱ Прогресс: {elapsed:.1f} мин | Итерация: {iteration}")

            time.sleep(POLL_INTERVAL_SEC)

    except Exception as e:
        print(f"[FATAL] Unexpected error: {e}")
        import traceback
        traceback.print_exc()
    finally:
        _graceful_shutdown()


if __name__ == "__main__":
    main()